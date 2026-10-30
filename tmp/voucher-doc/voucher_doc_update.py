from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
import shutil

p=next(Path('DOCUMENT').glob('*CUSTOMER*'))
backup=Path('tmp/voucher-doc'); backup.mkdir(parents=True,exist_ok=True)
shutil.copy2(p,backup/'original.docx')
d=Document(p)
anchor=next(x for x in d.paragraphs if x.text=='Schema request và response')._p
def para(text,style='Normal'):
    x=d.add_paragraph(text,style); anchor.addprevious(x._p); return x
def table(headers,rows):
    t=d.add_table(rows=1,cols=len(headers)); t.style=d.tables[1].style
    for c,v in zip(t.rows[0].cells,headers): c.text=v
    for vals in rows:
        for c,v in zip(t.add_row().cells,vals): c.text=v
    for r in t.rows:
        for c in r.cells:
            for q in c.paragraphs:
                for run in q.runs: run.font.size=Pt(9)
    anchor.addprevious(t._tbl)
    para('')

items=[
('GET','applicable','Danh sách voucher khả dụng','ApiResponse<List<VoucherResponse>>',
'Query: orderAmount (double, mặc định 0.0), serviceType (String, mặc định ALL), userId (String, tùy chọn). Không có JSON body. Nên truyền số tiền thực tế và RIDE_HAILING hoặc FOOD_DELIVERY; ALL không có nghĩa là lấy mọi voucher riêng của từng dịch vụ.',
'Lọc ACTIVE, mức tiền tối thiểu, dịch vụ, quota người dùng và số lượt còn lại. Implementation chưa lọc startAt/expireAt trực tiếp tại đây; phải gọi apply để xác nhận mã còn hợp lệ. Không phân trang.'),
('POST','apply','Kiểm tra và tính thử giảm giá','ApiResponse<ApplyVoucherResponse>',
'JSON ApplyVoucherRequest: code (String, bắt buộc không trắng), orderAmount (double, > 0), serviceType (String, bắt buộc không trắng; dùng RIDE_HAILING hoặc FOOD_DELIVERY), userId (String, tùy chọn), orderId (String, tùy chọn; chưa dùng trong tính thử).',
'Chuẩn hóa code bằng trim và viết hoa. Kiểm tra tồn tại, ACTIVE, thời gian kích hoạt/hết hạn, tiền tối thiểu, dịch vụ, quota và số lượt còn. Chỉ tính thử, không giữ chỗ hoặc trừ lượt. Mã không hợp lệ vẫn HTTP 200 nhưng success=false và data.valid=false; phải đọc message.'),
('POST','commit','Xác nhận sử dụng voucher','ApiResponse<ApplyVoucherResponse>',
'JSON CommitVoucherRequest: code và userId (String, bắt buộc không trắng); orderId, serviceType (String, tùy chọn theo DTO); originalAmount, discountAmount, finalAmount (double, thiếu sẽ là 0; không có annotation kiểm tra số tiền). Nên gửi orderId của đơn/chuyến đã tạo để liên kết lịch sử và hoàn lượt.',
'@Valid chạy trước thân controller, vì vậy userId bắt buộc trong JSON dù controller có nhánh lấy từ X-User-Id. Thành công trừ stock, tăng quota/usedCount và lưu VoucherUsage trạng thái USED. Hiện chưa chống commit trùng theo orderId; không tự retry. Backend nhận các số tiền do client gửi, chưa tính lại và chưa kiểm tra đầy đủ điều kiện như apply. HTTP 200 có thể success=false; không coi đây là xác nhận thanh toán hay cập nhật giá của đơn/chuyến.'),
('POST','rollback','Hoàn lượt sử dụng voucher','ApiResponse<Void>',
'Query: code, userId, orderId (String, đều tùy chọn theo controller). Không có JSON body. Nên gửi orderId và code của lượt đã commit; nếu thiếu code có thể tìm theo orderId. Nếu không tìm được theo orderId, backend có thể tìm lượt USED mới nhất theo code và userId.',
'Hoàn stock/quota, giảm usedCount và chuyển lịch sử sang REFUNDED, ghi refundedAt khi tìm thấy bản ghi. Chỉ bỏ qua hoàn lặp khi tìm được bản ghi đã REFUNDED; không bảo đảm idempotency cho mọi tổ hợp tham số. Không xác định được code thì bỏ qua nhưng vẫn trả success=true. Một số luồng backend đã tự hoàn lượt; cần đối chiếu lịch sử trước khi gọi bổ sung.'),
('GET','my-usages','Lịch sử sử dụng voucher','ApiResponse<List<VoucherUsage>>',
'Query userId (String, tùy chọn). Không có JSON body. Nếu không truyền, controller dùng X-User-Id từ Gateway.',
'Trả danh sách theo usedAt giảm dần, không phân trang. Thiếu định danh trả HTTP 200, success=true, data=[] và message báo User ID không hợp lệ. Gồm cả USED và REFUNDED.')]

for i,(method,route,title,response,inputs,notes) in enumerate(items,60):
    row=d.tables[0].add_row()
    for c,v in zip(row.cells,[str(i),f'{method} /api/v1/pricing/vouchers/{route}',title,'200',response]): c.text=v
    para(f'{i} {title}','Heading2')
    para(f'{method} /api/v1/pricing/vouchers/{route}')
    para('Xác thực: Bearer token qua Gateway. Gateway chuyển X-User-Id; controller chưa có kiểm tra role CUSTOMER/DRIVER riêng. userId được gửi tường minh được ưu tiên hơn header; app chỉ dùng định danh của phiên hiện tại, backend chưa đối chiếu quyền sở hữu tại các method này.')
    para(f'HTTP thành công: 200. Response: {response}.')
    para('Đầu vào: '+inputs)
    para(notes)
    para('Đầu ra: envelope gồm success (boolean), message (String), data và timestamp (LocalDateTime); các trường null của envelope bị lược bỏ. Với apply/commit, success bằng data.valid. Với rollback, không gán data.')
    para('Nguồn: pricing-service / VoucherController.java, VoucherServiceImpl.java và DTO tương ứng.')

para('Schema dữ liệu voucher','Heading2')
table(['Đối tượng','Trường và kiểu dữ liệu'],[
('ApplyVoucherResponse','valid: boolean; message: String; originalAmount, discountAmount, finalAmount: double; voucher: VoucherResponse (có thể null khi không hợp lệ).'),
('VoucherResponse','code, title, description: String; discountType: PERCENTAGE hoặc FIXED_AMOUNT; discountValue: double; maxDiscountAmount: Double (có thể null); minOrderAmount: double.'),
('VoucherResponse tiếp','startAt, expireAt, createdAt: LocalDateTime, yyyy-MM-dd\'T\'HH:mm:ss; totalUsageLimit, usedCount, remainingCount, usageLimitPerUser: int; applicableService: ALL, RIDE_HAILING, FOOD_DELIVERY; status: ACTIVE hoặc INACTIVE.'),
('VoucherUsage','id: Long; voucherCode, userId, orderId, serviceType: String; originalAmount, discountAmount, finalAmount: double; status: USED hoặc REFUNDED; usedAt, refundedAt: LocalDateTime, yyyy-MM-dd\'T\'HH:mm:ss. orderId, serviceType, refundedAt có thể null.')])
para('Ví dụ gọi voucher','Heading2')
para('Các mã, ID và số tiền dưới đây chỉ là ví dụ; thay bằng dữ liệu thực tế. Gửi Authorization: Bearer <access_token>; POST có JSON dùng Content-Type: application/json.')
for s in ['GET /api/v1/pricing/vouchers/applicable?orderAmount=100000&serviceType=RIDE_HAILING', 'POST /api/v1/pricing/vouchers/apply\n{"code":"GIAM10","orderAmount":100000,"serviceType":"RIDE_HAILING"}', 'POST /api/v1/pricing/vouchers/commit\n{"code":"GIAM10","userId":"123","orderId":"456","serviceType":"RIDE_HAILING","originalAmount":100000,"discountAmount":10000,"finalAmount":90000}', 'POST /api/v1/pricing/vouchers/rollback?code=GIAM10&orderId=456', 'GET /api/v1/pricing/vouchers/my-usages']:
    para(s)
para('Luồng tích hợp: lấy danh sách → apply để xem trước → xác nhận đơn/chuyến và commit một lần với ID thực → xem my-usages; chỉ rollback khi cần hoàn lượt và đã kiểm tra luồng tự hoàn của backend. Driver không có endpoint voucher riêng. API quản trị voucher nằm dưới /api/v1/pricing/admin/**, yêu cầu ADMIN và ngoài phạm vi màn hình Customer/Driver.')
for x in d.paragraphs:
    if x.text.startswith('Không tìm thấy controller mobile cho'):
        for r in x.runs: r.text=r.text.replace(', mã giảm giá','')
d.save(p)
check=Document(p)
assert all(any(f'/api/v1/pricing/vouchers/{r}' in q.text for q in check.paragraphs) for _,r,*_ in items)
print('Updated:',p,'; tables:',len(check.tables))
