from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
import re, shutil, json
p=next(Path('DOCUMENT').glob('*CUSTOMER*'))
shutil.copy2(p,Path('tmp/voucher-doc/before-admin.docx'))
d=Document(p)
anchor=next(q for q in d.paragraphs if q.text=='Schema request và response')._p
def para(text,style='Normal'):
    q=d.add_paragraph(text,style); anchor.addprevious(q._p); return q
base='/api/v1/pricing/admin/vouchers'
items=[
('POST','','Tạo voucher','201','VoucherResponse','JSON CreateVoucherRequest đầy đủ theo schema bên dưới. Không có path/query.', 'Chuẩn hóa mã thành chữ hoa; từ chối mã trùng trong DB hoặc Redis. Lưu PostgreSQL và Redis, usedCount khởi tạo 0, trạng thái lưu ACTIVE; trạng thái trả về có thể UPCOMING theo ngày bắt đầu.'),
('PUT','/{code}','Chỉnh sửa voucher','200','VoucherResponse','Path code là mã hiện tại. JSON CreateVoucherRequest chứa toàn bộ thông tin mới, bao gồm code trong body.', 'Đây là cập nhật toàn bộ, không phải PATCH từng trường. Có thể đổi mã bằng body.code khác path.code; mã mới không được trùng. Các trường tùy chọn bị bỏ qua trong JSON có thể trở về giá trị mặc định/null và ghi đè dữ liệu cũ. Giữ nguyên usedCount, trạng thái lưu và createdAt; cập nhật updatedAt, DB và Redis; xóa cache mã cũ khi đổi mã. Với FIXED_AMOUNT, cập nhật maxDiscountAmount thành null. Không có thao tác đổi voucherCode trong lịch sử hoặc di chuyển quota theo người dùng khi đổi mã.'),
('GET','','Danh sách voucher quản trị','200','List<VoucherResponse>','Không có body, path hoặc query; chưa có tìm kiếm, lọc hoặc phân trang.', 'Ưu tiên danh sách Redis. Chỉ khi Redis rỗng mới đọc toàn bộ DB và nạp cache các mã chưa hết hạn. Vì vậy danh sách không bảo đảm chứa toàn bộ voucher trong DB khi cache có một phần dữ liệu; thứ tự không được cam kết.'),
('GET','/{code}','Chi tiết voucher theo mã','200','VoucherResponse','Path code: String, mã cần tra cứu. Không có body/query.', 'Chuẩn hóa mã bằng trim và viết hoa; tìm Redis trước, fallback DB. Không tìm thấy ném ResourceNotFoundException. Trạng thái response được tính lại theo thời gian.'),
('PATCH','/{code}/toggle','Bật tắt voucher','200','VoucherResponse','Path code: String. Không có JSON body hoặc tham số status.', 'Đảo trạng thái lưu: ACTIVE thành INACTIVE, các trạng thái khác thành ACTIVE; đồng bộ DB/Redis. Không nhận status đích từ client, gọi lặp có thể đảo lại nên không tự retry. Response vẫn có thể EXPIRED hoặc UPCOMING theo thời gian; bật lại không tự gia hạn.'),
('DELETE','/{code}','Xóa voucher','200','Void','Path code: String. Không có body/query.', 'Xóa voucher khỏi PostgreSQL và Redis, không chỉ Redis như message controller. Không xóa lịch sử voucher_usages trong method này. Response có success, message, timestamp và không gán data. Không tìm thấy mã ném ResourceNotFoundException.'),
('GET','/{code}/usages','Lịch sử sử dụng theo voucher','200','List<VoucherUsage>','Path code: String. Không có body/query và không phân trang.', 'Chuẩn hóa mã, trả tất cả lượt dùng của mã theo usedAt giảm dần, gồm USED và REFUNDED. Không kiểm tra voucher tồn tại trước khi truy vấn; không có lịch sử trả data=[]. Schema VoucherUsage nằm trong phần dữ liệu voucher phía trên.')]
para('Quản trị voucher','Heading1')
para('Nhóm bổ sung dành cho ADMIN gồm tạo, sửa, danh sách, chi tiết, bật tắt, xóa và lịch sử theo mã. Cùng với 5 endpoint sử dụng voucher phía trên, tổng cộng có 12 endpoint voucher trong hai controller hiện tại. Các endpoint quản trị được ghi vào tài liệu để bao quát đầy đủ chức năng voucher; Customer và Driver không có quyền gọi nhóm này.')
para('Xác thực chung: Authorization: Bearer <access_token_admin>. Gateway chuyển định danh/quyền xuống service; SecurityConfig yêu cầu hasRole("ADMIN") cho /api/v1/pricing/admin/**. Body POST/PUT dùng Content-Type: application/json.')
for n,(method,suffix,title,status,typ,inputs,notes) in enumerate(items,65):
    values=[str(n),method+' '+base+suffix,title+' (ADMIN)',status,'ApiResponse<'+typ+'>']
    template=deepcopy(d.tables[0].rows[-1]._tr);d.tables[0]._tbl.append(template)
    row=d.tables[0].rows[-1]
    for c,v in zip(row.cells,values):
        c.text=v
    para(str(n)+' '+title,'Heading2')
    para(method+' '+base+suffix)
    para('Quyền: ADMIN. HTTP thành công: '+status+'. Response: ApiResponse<'+typ+'>.')
    para('Đầu vào: '+inputs)
    para(notes)
    para('Đầu ra: success=true, message, timestamp và data kiểu '+typ+' (Void không gán data). VoucherResponse và VoucherUsage dùng schema đã liệt kê trong phần voucher; các trường null của envelope bị lược bỏ.')
    para('Nguồn: pricing-service / AdminVoucherController.java, VoucherServiceImpl.java, SecurityConfig.java.')
para('Schema tạo và chỉnh sửa voucher','Heading2')
para('POST và PUT đều dùng CreateVoucherRequest với @Valid. PUT yêu cầu gửi lại các trường bắt buộc; không có DTO UpdateVoucherRequest riêng và không có endpoint PATCH /{code} để sửa một phần.')
fields=[
('code','String','Bắt buộc; regex ^[A-Za-z0-9_-]{3,20}$, gồm chữ, số, dấu gạch ngang hoặc gạch dưới; 3–20 ký tự. Body không nhận khoảng trắng thừa do validation chạy trước chuẩn hóa.'),
('title','String','Bắt buộc, không trắng.'),('description','String','Tùy chọn; có thể null.'),
('discountType','enum','Bắt buộc: PERCENTAGE hoặc FIXED_AMOUNT.'),('discountValue','double','Bắt buộc >0; với PERCENTAGE phải <=100, có thể là số thập phân.'),
('maxDiscountAmount','Double','Nếu có phải >0. Bắt buộc với PERCENTAGE; khi PUT FIXED_AMOUNT backend đặt null.'),('minOrderAmount','double','>=0, mặc định 0.'),
('startAt','LocalDateTime','Bắt buộc; yyyy-MM-dd\'T\'HH:mm:ss, không chứa offset.'),('expireAt','LocalDateTime','Bắt buộc; cùng định dạng startAt. Service từ chối khi trước startAt hoặc trước thời điểm hiện tại; phép so sánh hiện tại dùng isBefore nên không từ chối chỉ vì bằng startAt.'),
('totalUsageLimit','Integer','>=0; mặc định 0. Null hoặc 0 được lưu 0, nghĩa là không giới hạn tổng lượt.'),('usageLimitPerUser','int','>=1, mặc định 1.'),
('applicableService','enum','ALL, RIDE_HAILING hoặc FOOD_DELIVERY; mặc định ALL khi bỏ trường. PUT chuyển null tường minh thành ALL; luồng tạo không chuẩn hóa null tường minh, nên gửi ALL hoặc bỏ trường.')]
t=d.add_table(rows=1,cols=3);t.style=d.tables[1].style
for c,v in zip(t.rows[0].cells,['Trường','Kiểu','Ràng buộc và mặc định']): c.text=v
for values in fields:
    for c,v in zip(t.add_row().cells,values):c.text=v
for row in t.rows:
    for c in row.cells:
        for q in c.paragraphs:
            for r in q.runs:r.font.size=Pt(9)
anchor.addprevious(t._tbl)
para('Ví dụ chỉnh sửa voucher','Heading2')
para('PUT /api/v1/pricing/admin/vouchers/GIAM10')
example=dict(code='GIAM10',title='Giảm 15 phần trăm cho chuyến xe',description='Ưu đãi đặt xe',discountType='PERCENTAGE',discountValue=15,maxDiscountAmount=30000,minOrderAmount=50000,startAt='2026-11-01T00:00:00',expireAt='2026-12-31T23:59:59',totalUsageLimit=1000,usageLimitPerUser=2,applicableService='RIDE_HAILING')
para(json.dumps(example,ensure_ascii=False,indent=2))
para('Ví dụ minh họa: thay mã và thời gian theo dữ liệu thực tế. Để tạo mới, gửi cùng cấu trúc JSON tới POST /api/v1/pricing/admin/vouchers với code chưa tồn tại. Để đổi mã, giữ path là mã cũ và đổi code trong body; lưu ý lịch sử vẫn gắn mã cũ.')
para('Lỗi và lưu ý quản trị','Heading2')
para('Validation DTO hoặc JSON sai kiểu/enum có thể bị từ chối trước khi vào controller. Trùng mã, thời gian không hợp lệ hoặc điều kiện giảm giá sai ném BadRequestException; mã không tồn tại khi sửa/đọc/bật tắt/xóa ném ResourceNotFoundException. Hai exception hiện chỉ kế thừa RuntimeException, không có @ResponseStatus; không cam kết chúng được ánh xạ HTTP 400/404 hoặc envelope ApiResponse khi chưa có handler tương ứng. App phải xử lý cả response lỗi không theo envelope thành công.')
para('PUT không có kiểm tra totalUsageLimit mới phải lớn hơn usedCount, không tự thay đổi trạng thái lưu và không cập nhật lịch sử khi đổi mã. Nên đọc chi tiết trước khi sửa để giữ đủ các trường cần bảo toàn. Schema response dùng chung: VoucherResponse ở danh sách/chi tiết/tạo/sửa/bật tắt; VoucherUsage ở lịch sử.')
for q in d.paragraphs:
    if q.text=='Danh mục endpoint mobile':q.text='Danh mục endpoint mobile và quản trị voucher'
    if 'API quản trị voucher nằm dưới' in q.text:q.text=q.text.replace('và ngoài phạm vi màn hình Customer/Driver.','và được mô tả đầy đủ trong phần Quản trị voucher bên dưới.')
d.save(p)
# Compare all mappings in both source controllers against catalog and detail text.
catalog={r.cells[1].text for r in d.tables[0].rows[1:]}
details={q.text for q in d.paragraphs}
routes=[]
root=Path('SRC/BE/pricing-service/src/main/java/com/trung/pricingservice/controller')
for filename in ['VoucherController.java','AdminVoucherController.java']:
    source=(root/filename).read_text(encoding='utf-8')
    prefix=re.search(r'@RequestMapping\("([^"]+)"\)',source).group(1)
    for match in re.finditer(r'@(Get|Post|Put|Patch|Delete)Mapping(?:\("([^"]*)"\))?',source):
        route=match.group(1).upper()+' '+prefix+(match.group(2) or '')
        assert route in catalog and route in details,route
        routes.append(route)
assert len(routes)==12
Document(p)
print('Verified all 12 voucher endpoint mappings in catalog and details. DOCX reopened successfully.')
